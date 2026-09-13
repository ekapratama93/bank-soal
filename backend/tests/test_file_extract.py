import pytest

from app.file_extract import FileExtractError, extract_text


class TestFileExtract:
    def test_txt_utf8(self):
        text = extract_text("materi.txt", "Perkalian dasar.\n2 x 3 = 6".encode("utf-8"))
        assert "Perkalian dasar." in text

    def test_txt_non_utf8_fallback(self):
        text = extract_text("materi.txt", "Jaring-jaring kubus \xe9".encode("latin-1"))
        assert "Jaring-jaring" in text

    def test_txt_utf8_arabic(self):
        text = extract_text(
            "materi.txt",
            "Hadis: إنَّمَا الأَعْمَالُ بِالنِّيَّاتِ".encode("utf-8"),
        )
        assert "إنَّمَا" in text

    def test_txt_utf8_bom_arabic(self):
        text = extract_text(
            "materi.txt",
            "بسم الله".encode("utf-8-sig"),
        )
        assert text.startswith("بسم")

    def test_txt_cp1256_arabic(self):
        text = extract_text(
            "materi.txt",
            "بسم الله".encode("cp1256"),
        )
        assert text == "بسم الله"

    def test_pdf(self, monkeypatch):
        import app.file_extract as fe

        class FakePage:
            def extract_text(self):
                return "Isi halaman PDF"

        class FakeReader:
            def __init__(self, data):
                pass

            pages = [FakePage()]

        import sys

        monkeypatch.setitem(sys.modules, "pypdf", type(sys)("pypdf"))
        import pypdf

        pypdf.PdfReader = FakeReader
        text = extract_text("materi.pdf", b"%PDF-fake")
        assert "Isi halaman" in text

    def test_pdf_read_error(self, monkeypatch):
        import sys

        monkeypatch.setitem(sys.modules, "pypdf", type(sys)("pypdf"))
        import pypdf

        class FakeReader:
            def __init__(self, data):
                raise ValueError("corrupt")

        pypdf.PdfReader = FakeReader
        with pytest.raises(FileExtractError, match="PDF"):
            extract_text("materi.pdf", b"%PDF-fake")

    def test_docx(self):
        import io

        from docx import Document

        doc = Document()
        doc.add_paragraph("Bab 1: Perkalian")
        doc.add_paragraph("")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "2 x 3"
        table.cell(0, 1).text = "6"
        table.cell(1, 0).text = "3 x 4"
        table.cell(1, 1).text = "12"
        buf = io.BytesIO()
        doc.save(buf)
        text = extract_text("materi.docx", buf.getvalue())
        assert "Bab 1: Perkalian" in text
        assert "2 x 3\t6" in text

    def test_docx_corrupt(self):
        with pytest.raises(FileExtractError, match="DOCX"):
            extract_text("materi.docx", b"bukan file docx")

    def test_docx_empty(self):
        import io

        from docx import Document

        buf = io.BytesIO()
        Document().save(buf)
        with pytest.raises(FileExtractError, match="kosong"):
            extract_text("materi.docx", buf.getvalue())

    def test_unsupported_extension(self):
        with pytest.raises(FileExtractError, match=".pdf, .docx, atau .txt"):
            extract_text("materi.rtf", b"isi")

    def test_empty_text(self):
        with pytest.raises(FileExtractError, match="kosong"):
            extract_text("materi.txt", b"   \n  ")

    def test_size_limit(self):
        with pytest.raises(FileExtractError, match="2 MB"):
            extract_text("materi.txt", b"x" * (2 * 1024 * 1024 + 1))