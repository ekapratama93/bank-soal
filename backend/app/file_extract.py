import io

MAX_FILE_BYTES = 2 * 1024 * 1024  # 2 MB


class FileExtractError(Exception):
    pass


_TXT_ENCODINGS = ("utf-8-sig", "utf-8", "cp1256")


def _extract_txt(data: bytes) -> str:
    for encoding in _TXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileExtractError("File teks tidak dapat dibaca.")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages).strip()
    except FileExtractError:
        raise
    except Exception as e:
        raise FileExtractError("Gagal membaca file PDF.") from e


def _extract_docx(data: bytes) -> str:
    try:
        import docx

        doc = docx.Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                lines.append("\t".join(cells))
        return "\n".join(lines).strip()
    except FileExtractError:
        raise
    except Exception as e:
        raise FileExtractError("Gagal membaca file DOCX.") from e


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_FILE_BYTES:
        raise FileExtractError("Ukuran file maksimal 2 MB.")
    name = filename.lower()
    if name.endswith(".txt"):
        text = _extract_txt(data)
    elif name.endswith(".pdf"):
        text = _extract_pdf(data)
    elif name.endswith(".docx"):
        text = _extract_docx(data)
    else:
        raise FileExtractError("Format file harus .pdf, .docx, atau .txt")
    if not text.strip():
        raise FileExtractError(
            "Isi file kosong atau tidak terbaca. PDF hasil scan tidak didukung."
        )
    return text